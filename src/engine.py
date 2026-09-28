#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
L的工具箱 —— PowerPoint 图片标定、比例尺与原像素批量裁取引擎。

标定数据绑定到具体图片；图片移动或缩放后仍按原始像素换算。
批量裁取读取 PPT 内嵌原图，只新增裁出的小图，不改动原图片。

版本号的唯一来源（build/tools 从本行读取，不要在其他文件重复写死）。
"""
VERSION = "0.1.0"

import sys, os, math, time, argparse, contextlib, tempfile, io

# ---------- PowerPoint 常量 ----------
msoShapeRectangle = 1
msoAutoShape = 1
msoLine = 9
ppPicture = 13
ppAlignLeft, ppAlignCenter, ppAlignRight = 1, 2, 3
msoTrue, msoFalse = -1, 0

PIC_CAL_UM_PX_TAG = "LTOOLBOX_UM_PER_PX"
PIC_CAL_REAL_TAG = "LTOOLBOX_REAL_UM"
PIC_CAL_SPAN_PX_TAG = "LTOOLBOX_SPAN_PX"
BOX_PIC_ID_TAG = "LTOOLBOX_PICTURE_ID"
CROP_SOURCE_PIC_TAG = "LTOOLBOX_SOURCE_PICTURE_ID"
CROP_SOURCE_BOX_TAG = "LTOOLBOX_SOURCE_BOX_ID"
CROP_PIXEL_RECT_TAG = "LTOOLBOX_SOURCE_PIXEL_RECT"


# ---------- 连接 PowerPoint ----------
def get_pp():
    import win32com.client
    return win32com.client.GetActiveObject("PowerPoint.Application")


def active():
    pp = get_pp()
    pres = pp.ActivePresentation
    slide = pp.ActiveWindow.View.Slide
    return pp, pres, slide


def sel_shapes(slide):
    """返回当前选中的 shape 列表（可能多选）。"""
    pp = get_pp()
    sel = pp.ActiveWindow.Selection
    out = []
    if sel.Type == 2:  # ppSelectionShapes
        try:
            sr = sel.ShapeRange
            for i in range(1, sr.Count + 1):
                out.append(sr.Item(i))
        except Exception:
            pass
    return out






# ---------- 图片绑定标定 ----------
def _tag_value(owner, name):
    try:
        return owner.Tags.Item(name) or None
    except Exception:
        return None


def _tag_float(owner, name):
    value = _tag_value(owner, name)
    try:
        return float(value) if value is not None else None
    except (TypeError, ValueError):
        return None


def _set_tag(owner, name, value):
    owner.Tags.Add(name, str(value))


def _all_pictures(slide):
    return [slide.Shapes.Item(i) for i in range(1, slide.Shapes.Count + 1)
            if slide.Shapes.Item(i).Type == ppPicture]


def _contains(outer, inner, tolerance=2.0):
    return (inner.Left >= outer.Left - tolerance
            and inner.Top >= outer.Top - tolerance
            and inner.Left + inner.Width <= outer.Left + outer.Width + tolerance
            and inner.Top + inner.Height <= outer.Top + outer.Height + tolerance)


@contextlib.contextmanager
def _snapshot_deck(pres):
    from pptx import Presentation
    with tempfile.TemporaryDirectory(prefix="LToolbox_") as tmp:
        path = os.path.join(tmp, "snapshot.pptx")
        pres.SaveCopyAs(path, 24)
        yield Presentation(path)


def _picture_info(deck, slide, pic):
    from PIL import Image, ImageOps
    pptx_slide = deck.slides[slide.SlideIndex - 1]
    shape = next((s for s in pptx_slide.shapes
                  if s.shape_id == pic.Id and hasattr(s, "image")), None)
    if shape is None:
        raise SystemExit("无法读取图片“%s”的内嵌原图。" % pic.Name)
    blob = shape.image.blob
    try:
        with Image.open(io.BytesIO(blob)) as opened:
            image = ImageOps.exif_transpose(opened)
            size = image.size
    except Exception:
        raise SystemExit("图片“%s”不是可裁取的 PNG/JPEG/TIFF 位图。" % pic.Name)
    left = float(shape.crop_left or 0)
    right = float(shape.crop_right or 0)
    top = float(shape.crop_top or 0)
    bottom = float(shape.crop_bottom or 0)
    visible_w = size[0] * (1 - left - right)
    visible_h = size[1] * (1 - top - bottom)
    if visible_w <= 0 or visible_h <= 0:
        raise SystemExit("图片“%s”的 PowerPoint 裁剪参数无效。" % pic.Name)
    return {
        "blob": blob, "size": size,
        "x0": size[0] * left, "y0": size[1] * top,
        "visible_w": visible_w, "visible_h": visible_h,
        "px_per_pt_x": visible_w / pic.Width,
        "px_per_pt_y": visible_h / pic.Height,
    }


def get_calibration_record(pic):
    um_per_px = _tag_float(pic, PIC_CAL_UM_PX_TAG)
    if not um_per_px or um_per_px <= 0:
        return None
    return {
        "um_per_px": um_per_px,
        "real_um": _tag_float(pic, PIC_CAL_REAL_TAG),
        "span_px": _tag_float(pic, PIC_CAL_SPAN_PX_TAG),
    }


def _require_picture_calibration(pic):
    record = get_calibration_record(pic)
    if record is None:
        raise SystemExit(
            "图片“%s”尚未标定。\n\n请在这张图片自带比例尺上画等长直线，"
            "选中直线后点击【① 标定 / 状态】。" % pic.Name)
    return record


def _pick_picture(slide):
    selected = [s for s in sel_shapes(slide) if s.Type == ppPicture]
    if len(selected) == 1:
        return selected[0]
    if len(selected) > 1:
        raise SystemExit("一次只能为一张图片生成物理尺寸取景框，请只选中目标图片。")
    pictures = _all_pictures(slide)
    if len(pictures) == 1:
        return pictures[0]
    raise SystemExit("当前页有多张图片，请先选中目标图片。")


def _picture_for_line(slide, line):
    selected = [s for s in sel_shapes(slide) if s.Type == ppPicture]
    if len(selected) == 1:
        candidates = selected
    elif len(selected) > 1:
        raise SystemExit("标定时只能选中一张目标图片。")
    else:
        candidates = [p for p in _all_pictures(slide) if _contains(p, line)]
    if len(candidates) != 1:
        raise SystemExit("无法唯一确定标定线属于哪张图片；请同时选中标定线和目标图片。")
    if not _contains(candidates[0], line):
        raise SystemExit("标定线必须完整位于目标图片内。")
    return candidates[0]


def _status_badge(slide, pic, record):
    name = "LT_calibration_status_%s" % pic.Id
    for i in range(slide.Shapes.Count, 0, -1):
        shape = slide.Shapes.Item(i)
        if shape.Name in ("LT_calibration_status", name):
            shape.Delete()
    height = 18.0
    slide_width = get_pp().ActivePresentation.PageSetup.SlideWidth
    right_space = slide_width - pic.Left - pic.Width
    if right_space >= 130:
        width = min(180.0, right_space - 4)
        x, y = pic.Left + pic.Width + 4, pic.Top
        alignment = ppAlignLeft
    elif pic.Left >= 130:
        width = min(180.0, pic.Left - 4)
        x, y = pic.Left - width - 4, pic.Top
        alignment = ppAlignRight
    else:
        width = min(180.0, max(125.0, pic.Width * 0.34))
        x = pic.Left + pic.Width - width
        y = pic.Top - height if pic.Top >= height else pic.Top + 2
        alignment = ppAlignRight
    label = slide.Shapes.AddTextbox(1, x, y, width, height)
    label.Name = name
    ok = record is not None
    label.TextFrame.TextRange.Text = (
        "✓ 此图片已标定：%g µm" % record["real_um"]
        if ok else "✕ 此图片尚未标定")
    label.TextFrame.TextRange.Font.Size = 10
    label.TextFrame.TextRange.Font.Bold = msoTrue
    label.TextFrame.TextRange.Font.Color.RGB = 32768 if ok else 255
    label.TextFrame.TextRange.ParagraphFormat.Alignment = alignment
    label.Fill.Visible = msoTrue
    label.Fill.ForeColor.RGB = 16777215
    label.Fill.Transparency = 0.15
    label.Line.Visible = msoFalse
    label.Shadow.Visible = msoFalse
    _set_tag(label, BOX_PIC_ID_TAG, pic.Id)
    label.ZOrder(0)
    return label


def calibration_status():
    _, _, slide = active()
    pic = _pick_picture(slide)
    record = get_calibration_record(pic)
    _status_badge(slide, pic, record)
    if record is None:
        print("图片“%s”：未标定。请选择这张图上的标定线后再次点击①。" % pic.Name)
        return None
    print("图片“%s”：已标定。\n%g µm = %.1f 原图像素\n换算率 %.6g µm/px"
          % (pic.Name, record["real_um"], record["span_px"],
             record["um_per_px"]))
    return record


def calibrate(real_um, unit="um"):
    """用原图像素长度标定标尺所在的具体图片。"""
    _, pres, slide = active()
    line = next((s for s in sel_shapes(slide) if s.Type == msoLine), None)
    if line is None:
        raise SystemExit("请先在原图自带比例尺上画等长直线，并选中该直线。")
    pic = _picture_for_line(slide, line)
    if real_um <= 0:
        raise SystemExit("比例尺真实长度必须大于 0。")
    unit_factor = {"um": 1.0, "mm": 1000.0, "nm": 1e-3, "cm": 1e4}[unit]
    real_um_val = real_um * unit_factor
    with _snapshot_deck(pres) as deck:
        info = _picture_info(deck, slide, pic)
    try:
        dx = (line.EndX - line.BeginX) * info["px_per_pt_x"]
        dy = (line.EndY - line.BeginY) * info["px_per_pt_y"]
    except Exception:
        dx = line.Width * info["px_per_pt_x"]
        dy = line.Height * info["px_per_pt_y"]
    span_px = math.hypot(dx, dy)
    if span_px < 2:
        raise SystemExit("标定线太短，请重新画一条与原图比例尺等长的直线。")
    um_per_px = real_um_val / span_px
    _set_tag(pic, PIC_CAL_UM_PX_TAG, "%.12g" % um_per_px)
    _set_tag(pic, PIC_CAL_REAL_TAG, "%.12g" % real_um_val)
    _set_tag(pic, PIC_CAL_SPAN_PX_TAG, "%.12g" % span_px)
    for legacy_tag in (
            "LTOOLBOX_UM_PER_PT", "LTOOLBOX_REAL_UM", "LTOOLBOX_SPAN_PT"):
        try:
            slide.Tags.Delete(legacy_tag)
        except Exception:
            pass
    _set_tag(line, BOX_PIC_ID_TAG, pic.Id)
    line.Name = "LT_calibration_line_%s" % pic.Id
    line.Line.ForeColor.RGB = 32768
    line.Line.Weight = 2.25
    line.Shadow.Visible = msoFalse
    record = get_calibration_record(pic)
    _status_badge(slide, pic, record)
    print("图片“%s”标定完成。\n%g µm = %.1f 原图像素\n"
          "该图片复制、缩放或裁取后会继承标定。"
          % (pic.Name, real_um_val, span_px))
    return record






# ---------- 生成 µm 目标框 ----------
def make_box(target_um):
    """按所选图片的像素标定生成物理尺寸取景框；旧框保留。"""
    _, pres, slide = active()
    if target_um <= 0:
        raise SystemExit("取景框实际边长必须大于 0 µm。")
    pic = _pick_picture(slide)
    record = _require_picture_calibration(pic)
    with _snapshot_deck(pres) as deck:
        info = _picture_info(deck, slide, pic)
    um_per_pt_x = record["um_per_px"] * info["px_per_pt_x"]
    um_per_pt_y = record["um_per_px"] * info["px_per_pt_y"]
    width_pt = target_um / um_per_pt_x
    height_pt = target_um / um_per_pt_y
    if width_pt > pic.Width or height_pt > pic.Height:
        raise SystemExit(
            "目标取景框 %.1f µm 超出图片视野；当前图片最大可用边长约 %.1f µm。"
            % (target_um, min(pic.Width * um_per_pt_x,
                             pic.Height * um_per_pt_y)))
    old_boxes = [slide.Shapes.Item(i) for i in range(1, slide.Shapes.Count + 1)
                 if slide.Shapes.Item(i).Type == msoAutoShape
                 and slide.Shapes.Item(i).Name.startswith("LT_box")
                 and _contains(pic, slide.Shapes.Item(i))]
    recs = [{"x": b.Left, "y": b.Top, "w": b.Width, "h": b.Height}
            for b in old_boxes]
    left = pic.Left + (pic.Width - width_pt) / 2.0
    top = pic.Top + (pic.Height - height_pt) / 2.0
    left, top, moved = _avoid_overlap(
        left, top, width_pt, height_pt, recs,
        bounds=(pic.Left, pic.Top, pic.Width, pic.Height))
    if any(iou((left, top, width_pt, height_pt),
               (r["x"], r["y"], r["w"], r["h"])) > 0.01 for r in recs):
        raise SystemExit("当前图片没有足够的不重叠区域，请手动移动现有红框后重试。")
    used = {slide.Shapes.Item(i).Name
            for i in range(1, slide.Shapes.Count + 1)}
    number = 1
    while "LT_box_%02d" % number in used:
        number += 1
    box = slide.Shapes.AddShape(
        msoShapeRectangle, left, top, width_pt, height_pt)
    box.Name = "LT_box_%02d" % number
    _set_tag(box, BOX_PIC_ID_TAG, pic.Id)
    box.Fill.Visible = msoFalse
    box.Line.Weight = 3
    box.Line.ForeColor.RGB = 255
    box.Shadow.Visible = msoFalse
    box.ZOrder(0)
    _status_badge(slide, pic, record)
    print("图片“%s”已生成 %.1f × %.1f µm 新取景框%s。\n\n"
          "旧红框保留；需要原像素小图时，选中图片和一个或多个框后点击【批量原图裁取】。"
          % (pic.Name, target_um, target_um,
             "（已自动避开旧框）" if moved else ""))
    return box


def _avoid_overlap(left, top, w, h, recs, bounds=None):
    """在图片范围内寻找与旧红框基本不相交的位置。"""
    moved = False
    bx = by = 0.0
    bw = bh = 1e9
    if bounds:
        bx, by, bw, bh = bounds
    for _ in range(128):
        if not any(iou((left, top, w, h),
                       (r["x"], r["y"], r["w"], r["h"])) > 0.01
                   for r in recs):
            break
        left += w * 1.05
        if left + w > bx + bw:
            left = bx
            top += h * 1.05
        if top + h > by + bh:
            top = by
        moved = True
    left = min(max(left, bx), bx + bw - w)
    top = min(max(top, by), by + bh - h)
    return left, top, moved


def iou(a, b):
    ax, ay, aw, ah = a; bx, by, bw, bh = b
    x1 = max(ax, bx); y1 = max(ay, by)
    x2 = min(ax + aw, bx + bw); y2 = min(ay + ah, by + bh)
    iw = max(0, x2 - x1); ih = max(0, y2 - y1)
    inter = iw * ih
    union = aw * ah + bw * bh - inter
    return inter / union if union > 0 else 0.0


# ---------- 在取景框内生成比例尺 ----------
def _nice_scale_um(view_um):
    """选最接近图宽 20%、且不超过 25% 的 1/2/5×10ⁿ 比例尺。"""
    if view_um <= 0:
        raise ValueError("view_um must be positive")
    candidates = []
    for exponent in range(-3, 7):
        candidates.extend(m * (10 ** exponent) for m in (1, 2, 5))
    valid = [v for v in candidates if v <= view_um * 0.250001]
    if not valid:
        return min(candidates)
    target = view_um * 0.20
    return min(valid, key=lambda v: abs(math.log(v / target)))


def _current_box(slide):
    selected = [s for s in sel_shapes(slide)
                if s.Type == msoAutoShape
                and not s.Name.startswith("LT_scalebar")]
    if selected:
        return selected[-1]
    boxes = [slide.Shapes.Item(i) for i in range(1, slide.Shapes.Count + 1)
             if slide.Shapes.Item(i).Type == msoAutoShape
             and slide.Shapes.Item(i).Name.startswith("LT_box")]
    if not boxes:
        raise SystemExit("没有找到取景框。请先画矩形并选中，或点击【② 生成取景框】。")
    return boxes[-1]


def _picture_for_box(slide, box, pictures=None):
    pictures = pictures or _all_pictures(slide)
    owner = _tag_value(box, BOX_PIC_ID_TAG)
    if owner:
        owned = [p for p in pictures if str(p.Id) == str(owner)]
        if len(owned) == 1 and _contains(owned[0], box):
            return owned[0]
    candidates = [p for p in pictures if _contains(p, box)]
    if len(candidates) == 1:
        _set_tag(box, BOX_PIC_ID_TAG, candidates[0].Id)
        return candidates[0]
    if not candidates:
        raise SystemExit("取景框“%s”没有完整位于任何目标图片内。" % box.Name)
    raise SystemExit("取景框“%s”同时位于多张图片内；请只选中对应图片。" % box.Name)


def add_scalebar_to_box():
    """在当前框右下角生成规范比例尺；图片和红框均不改动。"""
    _, pres, slide = active()
    box = _current_box(slide)
    pic = _picture_for_box(slide, box)
    record = _require_picture_calibration(pic)
    with _snapshot_deck(pres) as deck:
        info = _picture_info(deck, slide, pic)
    um_per_pt_x = record["um_per_px"] * info["px_per_pt_x"]
    um_per_pt_y = record["um_per_px"] * info["px_per_pt_y"]
    view_w_um = box.Width * um_per_pt_x
    view_h_um = box.Height * um_per_pt_y
    scale_um = _nice_scale_um(view_w_um)
    bar_len = scale_um / um_per_pt_x
    pad = box.Width * 0.055
    thickness = max(2.0, min(5.0, box.Width * 0.018))
    font_size = max(8.0, min(14.0, box.Width * 0.055))
    label_width = max(bar_len, font_size * 4.5)
    right = box.Left + box.Width - pad
    x = right - bar_len
    label_x = right - label_width
    y = box.Top + box.Height - pad - thickness
    if box.Name.startswith("LT_box"):
        suffix = box.Name.removeprefix("LT_box").strip("_") or "01"
    else:
        suffix = str(box.Id)
    bar_name = "LT_scalebar_%s" % suffix
    label_name = "LT_scalelabel_%s" % suffix
    for i in range(slide.Shapes.Count, 0, -1):
        if slide.Shapes.Item(i).Name in (bar_name, label_name):
            slide.Shapes.Item(i).Delete()
    bar = slide.Shapes.AddShape(
        msoShapeRectangle, x, y, bar_len, thickness)
    bar.Name = bar_name
    _set_tag(bar, CROP_SOURCE_BOX_TAG, box.Id)
    bar.Fill.ForeColor.RGB = 0
    bar.Fill.Visible = msoTrue
    bar.Line.Visible = msoFalse
    bar.Shadow.Visible = msoFalse
    label = slide.Shapes.AddTextbox(
        1, label_x, y - font_size * 1.5, label_width, font_size * 1.4)
    label.Name = label_name
    _set_tag(label, CROP_SOURCE_BOX_TAG, box.Id)
    label.TextFrame.MarginLeft = 0
    label.TextFrame.MarginRight = 0
    label.TextFrame.MarginTop = 0
    label.TextFrame.MarginBottom = 0
    text = label.TextFrame.TextRange
    text.Text = "%g µm" % scale_um
    text.ParagraphFormat.Alignment = ppAlignRight
    text.Font.Size = font_size
    text.Font.Bold = msoTrue
    text.Font.Color.RGB = 0
    label.Fill.Visible = msoFalse
    label.Line.Visible = msoFalse
    label.Shadow.Visible = msoFalse
    bar.ZOrder(0)
    label.ZOrder(0)
    box.ZOrder(0)
    print("比例尺已生成在框内右下角，红框和原图保留。\n"
          "框内实际视野：%.1f × %.1f µm\n"
          "比例尺：%g µm（约占框宽 %.0f%%）"
          % (view_w_um, view_h_um, scale_um,
             100 * scale_um / view_w_um))
    return bar, label


# ---------- 原图像素批量裁取 ----------
def _crop_pixel_rect(info, pic, box):
    x1 = info["x0"] + (box.Left - pic.Left) * info["px_per_pt_x"]
    y1 = info["y0"] + (box.Top - pic.Top) * info["px_per_pt_y"]
    x2 = info["x0"] + (box.Left + box.Width - pic.Left) * info["px_per_pt_x"]
    y2 = info["y0"] + (box.Top + box.Height - pic.Top) * info["px_per_pt_y"]
    width, height = info["size"]
    return (
        max(0, min(width, math.floor(x1))),
        max(0, min(height, math.floor(y1))),
        max(0, min(width, math.ceil(x2))),
        max(0, min(height, math.ceil(y2))),
    )


def _bring_box_overlays_front(slide, box):
    overlays = [slide.Shapes.Item(i)
                for i in range(1, slide.Shapes.Count + 1)
                if _tag_value(slide.Shapes.Item(i), CROP_SOURCE_BOX_TAG)
                == str(box.Id)]
    box.ZOrder(0)
    for shape in overlays:
        shape.ZOrder(0)


def crop_selected_regions():
    """把选中框对应的原图像素裁成新图片，并叠放在各框原位置。"""
    from PIL import Image, ImageOps
    _, pres, slide = active()
    selected = sel_shapes(slide)
    pictures = [s for s in selected if s.Type == ppPicture]
    boxes = [s for s in selected
             if s.Type == msoAutoShape
             and not s.Name.startswith("LT_scalebar")]
    if not pictures or not boxes:
        raise SystemExit("请同时选中至少一张图片和一个矩形框，再点击【批量原图裁取】。")
    for pic in pictures:
        if abs(float(pic.Rotation or 0)) % 360 > 0.01:
            raise SystemExit("图片“%s”已旋转；请先恢复为 0° 后再原图裁取。" % pic.Name)
    assignments = [(box, _picture_for_box(slide, box, pictures))
                   for box in boxes]
    with _snapshot_deck(pres) as deck, tempfile.TemporaryDirectory(
            prefix="LToolbox_crops_") as tmp:
        infos = {pic.Id: _picture_info(deck, slide, pic) for pic in pictures}
        images = {}
        outputs = []
        details = []
        for number, (box, pic) in enumerate(assignments, 1):
            info = infos[pic.Id]
            if pic.Id not in images:
                with Image.open(io.BytesIO(info["blob"])) as opened:
                    images[pic.Id] = ImageOps.exif_transpose(opened).copy()
            pixel_rect = _crop_pixel_rect(info, pic, box)
            if pixel_rect[2] - pixel_rect[0] < 2 or pixel_rect[3] - pixel_rect[1] < 2:
                raise SystemExit("取景框“%s”对应的原图区域过小。" % box.Name)
            cropped = images[pic.Id].crop(pixel_rect)
            path = os.path.join(tmp, "crop_%03d.png" % number)
            try:
                cropped.save(path, "PNG")
            except OSError:
                cropped.convert("RGBA").save(path, "PNG")
            crop_name = "LT_crop_%s" % box.Id
            for i in range(slide.Shapes.Count, 0, -1):
                if slide.Shapes.Item(i).Name == crop_name:
                    slide.Shapes.Item(i).Delete()
            result = slide.Shapes.AddPicture(
                path, msoFalse, msoTrue,
                box.Left, box.Top, box.Width, box.Height)
            result.Name = crop_name
            _set_tag(result, CROP_SOURCE_PIC_TAG, pic.Id)
            _set_tag(result, CROP_SOURCE_BOX_TAG, box.Id)
            _set_tag(result, CROP_PIXEL_RECT_TAG, ",".join(map(str, pixel_rect)))
            record = get_calibration_record(pic)
            if record:
                _set_tag(result, PIC_CAL_UM_PX_TAG, "%.12g" % record["um_per_px"])
                _set_tag(result, PIC_CAL_REAL_TAG, "%.12g" % record["real_um"])
                _set_tag(result, PIC_CAL_SPAN_PX_TAG, "%.12g" % record["span_px"])
            result.AlternativeText = (
                "由图片“%s”按框“%s”从原图像素裁取；像素区域 %s。"
                % (pic.Name, box.Name, pixel_rect))
            _set_tag(box, BOX_PIC_ID_TAG, pic.Id)
            _bring_box_overlays_front(slide, box)
            outputs.append(result)
            details.append("%s：%d × %d px" % (
                box.Name, cropped.width, cropped.height))
        outputs[0].Select(msoTrue)
        for result in outputs[1:]:
            result.Select(msoFalse)
    print("已按原图分辨率裁出 %d 张小图，并放在各自框的位置；"
          "原图片和红框均未改动。\n%s"
          % (len(outputs), "\n".join(details)))
    return outputs


def _clear_notice():
    try:
        _, _, slide = active()
        for i in range(slide.Shapes.Count, 0, -1):
            if slide.Shapes.Item(i).Name == "LT_notice":
                slide.Shapes.Item(i).Delete()
    except Exception:
        pass


def _error_notice(message):
    try:
        _, pres, slide = active()
        _clear_notice()
        width = min(430.0, pres.PageSetup.SlideWidth - 24)
        notice = slide.Shapes.AddTextbox(
            1, (pres.PageSetup.SlideWidth - width) / 2, 8, width, 42)
        notice.Name = "LT_notice"
        notice.TextFrame.TextRange.Text = "L的工具箱：%s" % str(message)
        notice.TextFrame.TextRange.Font.Size = 11
        notice.TextFrame.TextRange.Font.Bold = msoTrue
        notice.TextFrame.TextRange.Font.Color.RGB = 255
        notice.Fill.Visible = msoTrue
        notice.Fill.ForeColor.RGB = 16777215
        notice.Line.Visible = msoTrue
        notice.Line.ForeColor.RGB = 255
        notice.ZOrder(0)
    except Exception:
        pass


# ---------- CLI ----------
def _msgbox(text, ok=True):
    """在 Windows 上弹置顶窗口（供功能区按钮调用时反馈结果）。"""
    try:
        import ctypes
        # 0x40=MB_ICONINFORMATION, 0x10=MB_ICONHAND, 0x40000=MB_TOPMOST, 0x10000=MB_SETFOREGROUND
        MB = (0x40 if ok else 0x10) | 0x40000 | 0x10000
        ctypes.windll.user32.MessageBoxW(0, str(text), "L的工具箱", MB)
    except Exception:
        print(text)


def _dispatch(a):
    buf = io.StringIO()
    _clear_notice()
    with contextlib.redirect_stdout(buf):
        if a.action == "status":
            calibration_status()
        elif a.action == "calibrate":
            calibrate(float(a.value), a.unit)
        elif a.action == "makebox":
            make_box(float(a.value))
        elif a.action == "scalebar":
            add_scalebar_to_box()
        elif a.action == "crop":
            crop_selected_regions()
        elif a.action == "version":
            print("L的工具箱 v%s" % VERSION)
    return buf.getvalue()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "action",
        choices=["status", "calibrate", "makebox", "scalebar", "crop", "version"])
    ap.add_argument("value", nargs="?", default=None)
    ap.add_argument("--unit", default="um")
    ap.add_argument("--gui", action="store_true", help="弹结果窗口（按钮调用）")
    # 打包成单文件 exe 时，VBA 启动串可能带一个空参数；空参数一律丢弃。
    a = ap.parse_args([x for x in sys.argv[1:] if x])
    log_p = os.path.join(tempfile.gettempdir(), "LToolbox_engine.log")
    try:
        with open(log_p, "a", encoding="utf-8") as f:
            f.write(f"[{time.strftime('%Y-%m-%d %H:%M:%S')}] sys.argv={sys.argv}\n")
    except Exception:
        pass
    if a.gui:
        time.sleep(0.5)  # 让 Ribbon VBA 回调先返回，避免 PowerPoint COM 忙碌
    try:
        out = _dispatch(a)
        with open(log_p, "a", encoding="utf-8") as f:
            f.write(f"Success out:\n{out}\n")
        if sys.stdout is not None:
            sys.stdout.write(out)
        if a.gui:
            _msgbox(out.strip() or "完成")
    except SystemExit as e:
        m = str(e)
        _error_notice(m)
        with open(log_p, "a", encoding="utf-8") as f:
            f.write(f"SystemExit: {m}\n")
        if sys.stderr is not None:
            sys.stderr.write(m + "\n")
        if a.gui:
            _msgbox(m, ok=False)
    except Exception as e:
        import traceback
        _error_notice("出错：" + str(e))
        tb = traceback.format_exc()
        with open(log_p, "a", encoding="utf-8") as f:
            f.write(f"Exception:\n{tb}\n")
        if sys.stderr is not None:
            sys.stderr.write(tb + "\n")
        if a.gui:
            _msgbox("出错\n" + str(e), ok=False)


if __name__ == "__main__":
    main()
